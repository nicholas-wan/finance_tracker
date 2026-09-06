"""Merge Trip.com "My Bookings" Excel exports into manual/trip_bookings.json.

Trip.com only lets you export the bookings that are currently on screen, so a
full history arrives as several partially overlapping workbooks. This script
reads them in the order given, keys every row on its booking number and keeps
one record per booking. A booking that appears in more than one export is
attributed to the last export that carried it, but it stays at the position
where it was first seen, so the merged file reads in export order.
"""

import argparse
import collections
import json
import os
import tempfile
from datetime import date
from pathlib import Path

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = REPO_ROOT / "manual" / "trip_bookings.json"
IDENTITY_PATH = REPO_ROOT / "manual" / "identity.json"

SHEET_NAME = "My Bookings"
SOURCE_LABEL = "Trip.com My Bookings Excel exports"
SCOPE = "Combined unique rows from all supplied exports, deduplicated by booking number"

# The workbook header, already stripped. The real exports ship "Booking No. "
# with a trailing space, so every header cell is stripped before comparing.
EXPECTED_HEADER = (
    "Booking status",
    "Product Type",
    "Booking No.",
    "Booking Date",
    "Product name",
    "Travel time",
    "Traveller",
    "Currency",
    "Amount",
)

# Worksheet column index -> output field name, in output key order.
COLUMNS = (
    (2, "bookingNo"),
    (0, "status"),
    (1, "productType"),
    (3, "bookingDate"),
    (4, "productName"),
    (5, "travelTime"),
    (6, "traveller"),
    (7, "currency"),
    (8, "amount"),
)
COMPARED_FIELDS = tuple(name for _, name in COLUMNS if name != "bookingNo")

# manual/trip_bookings.json is written with two-space indentation and real
# UTF-8 characters; keep matching it so re-imports produce no spurious diff.
JSON_INDENT = 2


def normalise_booking_no(value):
    """Render a booking number cell as a plain digit string."""
    if isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if value is None:
        return ""
    return str(value).strip()


def normalise_amount(value, path, row_number):
    if value is None:
        return None
    if isinstance(value, bool):
        raise SystemExit(
            "%s row %d: unsupported Amount value" % (path.name, row_number)
        )
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return round(value, 2)
    raise SystemExit(
        "%s row %d: Amount is not a number (%s)" % (path.name, row_number, type(value).__name__)
    )


def is_blank_row(row):
    return all(
        cell is None or (isinstance(cell, str) and not cell.strip()) for cell in row
    )


def read_workbook(path):
    """Read one export and return its rows as booking dicts, in sheet order."""
    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET_NAME in workbook.sheetnames:
            sheet = workbook[SHEET_NAME]
        else:
            sheet = workbook[workbook.sheetnames[0]]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows:
        raise SystemExit("%s: sheet is empty" % path.name)

    header = tuple(
        (cell.strip() if isinstance(cell, str) else cell) for cell in rows[0][: len(EXPECTED_HEADER)]
    )
    if header != EXPECTED_HEADER or len(rows[0]) != len(EXPECTED_HEADER):
        raise SystemExit(
            "%s: unexpected header row; expected %s" % (path.name, ", ".join(EXPECTED_HEADER))
        )

    bookings = []
    for offset, row in enumerate(rows[1:]):
        row_number = offset + 2  # 1-based, header is row 1
        if is_blank_row(row):
            continue
        record = {}
        for index, field in COLUMNS:
            value = row[index] if index < len(row) else None
            if field == "bookingNo":
                record[field] = normalise_booking_no(value)
            elif field == "amount":
                record[field] = normalise_amount(value, path, row_number)
            elif field == "currency":
                record[field] = ("" if value is None else str(value)).strip().upper()
            else:
                record[field] = ("" if value is None else str(value)).strip()
        if not record["bookingNo"]:
            raise SystemExit("%s row %d: missing booking number" % (path.name, row_number))
        record["sourceFile"] = path.name
        bookings.append(record)
    return bookings


def merge_bookings(rows_by_file, updates=None):
    """Collapse per-file rows onto one record per booking number.

    ``rows_by_file`` maps a source basename to its rows, in merge order. A
    booking seen again in a later file keeps its original position but takes
    that file's values and basename. ``updates``, when given, is a Counter that
    collects which fields a later export changed.
    """
    if hasattr(rows_by_file, "items"):
        ordered = list(rows_by_file.items())
    else:
        ordered = list(rows_by_file)

    merged = collections.OrderedDict()
    for _, rows in ordered:
        for record in rows:
            key = record["bookingNo"]
            previous = merged.get(key)
            if previous is not None and updates is not None:
                changed = [
                    field
                    for field in COMPARED_FIELDS
                    if previous.get(field) != record.get(field)
                ]
                if changed:
                    updates["__bookings__"] += 1
                    for field in changed:
                        updates[field] += 1
            merged[key] = dict(record)
    return list(merged.values())


def default_account_holder(identity_path=IDENTITY_PATH):
    path = Path(identity_path)
    if not path.exists():
        return ""
    try:
        with path.open(encoding="utf-8") as handle:
            identity = json.load(handle)
    except (OSError, ValueError):
        return ""
    value = identity.get("statementHolderName") if isinstance(identity, dict) else None
    return value if isinstance(value, str) else ""


def build_payload(bookings, source_files, account_holder="", imported_at=None):
    return {
        "source": SOURCE_LABEL,
        "sourceFiles": list(source_files),
        "accountHolder": account_holder,
        "importedAt": imported_at or date.today().isoformat(),
        "scope": SCOPE,
        "bookings": bookings,
    }


def render(payload):
    return json.dumps(payload, indent=JSON_INDENT, ensure_ascii=False) + "\n"


def write_atomic(path, text):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary = tempfile.mkstemp(dir=str(path.parent), prefix=path.name + ".", suffix=".tmp")
    try:
        with os.fdopen(handle, "w", encoding="utf-8", newline="\n") as stream:
            stream.write(text)
        os.replace(temporary, path)
    except BaseException:
        if os.path.exists(temporary):
            os.unlink(temporary)
        raise


def expand_inputs(arguments):
    """Turn CLI paths into a workbook list, expanding directories by name."""
    paths = []
    for argument in arguments:
        path = Path(argument)
        if path.is_dir():
            found = sorted(
                (child for child in path.glob("*.xlsx") if not child.name.startswith("~$")),
                key=lambda child: child.name,
            )
            if not found:
                raise SystemExit("%s: no .xlsx workbooks found" % path)
            paths.extend(found)
        elif path.exists():
            paths.append(path)
        else:
            raise SystemExit("%s: no such file or directory" % path)
    if not paths:
        raise SystemExit("no workbooks to read")
    return paths


def summarise(bookings):
    currencies = collections.Counter(record["currency"] for record in bookings)
    cancelled = sum(
        1 for record in bookings if record["status"].strip().lower() == "cancelled"
    )
    return currencies, cancelled


def compare_with_existing(output_path, payload):
    """Return a list of human-readable differences, ignoring volatile keys."""
    path = Path(output_path)
    if not path.exists():
        return ["%s does not exist" % path]
    try:
        with path.open(encoding="utf-8") as handle:
            existing = json.load(handle)
    except (OSError, ValueError) as error:
        return ["%s could not be read (%s)" % (path, error.__class__.__name__)]

    differences = []
    if existing.get("sourceFiles") != payload["sourceFiles"]:
        differences.append("sourceFiles differ")

    old = existing.get("bookings") or []
    new = payload["bookings"]
    if len(old) != len(new):
        differences.append("booking count %d vs %d" % (len(old), len(new)))

    old_by_no = {record.get("bookingNo"): record for record in old}
    new_by_no = {record["bookingNo"]: record for record in new}
    missing = set(old_by_no) - set(new_by_no)
    added = set(new_by_no) - set(old_by_no)
    if missing:
        differences.append("%d bookings only in the existing file" % len(missing))
    if added:
        differences.append("%d bookings only in the generated file" % len(added))

    changed = sum(
        1
        for key in set(old_by_no) & set(new_by_no)
        if old_by_no[key] != new_by_no[key]
    )
    if changed:
        differences.append("%d bookings with changed fields" % changed)

    old_order = [record.get("bookingNo") for record in old]
    new_order = [record["bookingNo"] for record in new]
    if not differences and old_order != new_order:
        differences.append("booking order differs")
    return differences


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "workbooks", nargs="+", metavar="WORKBOOK",
        help="Trip.com export .xlsx files, or directories of them, in merge order",
    )
    parser.add_argument(
        "--output", default=str(DEFAULT_OUTPUT),
        help="where to write the merged JSON (default: manual/trip_bookings.json)",
    )
    parser.add_argument(
        "--account-holder", default=None,
        help="account holder name (default: statementHolderName from manual/identity.json)",
    )
    parser.add_argument(
        "--imported-at", default=None,
        help="import date to stamp, YYYY-MM-DD (default: today)",
    )
    parser.add_argument(
        "--check", action="store_true",
        help="compare against the existing output instead of writing it",
    )
    args = parser.parse_args(argv)

    paths = expand_inputs(args.workbooks)
    rows_by_file = collections.OrderedDict()
    total_rows = 0
    for path in paths:
        rows = read_workbook(path)
        total_rows += len(rows)
        if path.name in rows_by_file:
            rows_by_file[path.name] = rows_by_file[path.name] + rows
        else:
            rows_by_file[path.name] = rows

    updates = collections.Counter()
    bookings = merge_bookings(rows_by_file, updates=updates)

    account_holder = args.account_holder
    if account_holder is None:
        account_holder = default_account_holder()
    payload = build_payload(
        bookings, list(rows_by_file), account_holder, args.imported_at
    )

    if args.check:
        differences = compare_with_existing(args.output, payload)
        if not differences:
            print("up to date")
            return 0
        print("%d difference(s): %s" % (len(differences), "; ".join(differences)))
        return 1

    write_atomic(args.output, render(payload))

    currencies, cancelled = summarise(bookings)
    print(
        "Read %d workbook(s), %d row(s) -> %d unique booking(s), %d duplicate(s) collapsed."
        % (len(paths), total_rows, len(bookings), total_rows - len(bookings))
    )
    if updates["__bookings__"]:
        detail = ", ".join(
            "%s: %d" % (field, updates[field])
            for field in COMPARED_FIELDS
            if updates[field]
        )
        print(
            "%d booking(s) updated by a later export (%s)."
            % (updates["__bookings__"], detail)
        )
    print(
        "Currencies: %s."
        % ", ".join("%s %d" % (code, count) for code, count in sorted(currencies.items()))
    )
    print("Cancelled bookings: %d." % cancelled)
    print("Wrote %s." % args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
