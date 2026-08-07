# One-time migration. Lifts everything hand-entered out of Finances.xlsx into
# manual/*.json so the spreadsheet is never needed again.
#
#   owner_tags.json   per-transaction Nic/Shared/Yx tags, keyed by month+desc+amount
#   owner_rules.json  merchant -> default owner, from the majority of history
#   salary.json       salary steps and annual income/tax
#
# Safe to re-run: it merges rather than replaces, so newer hand edits survive.

import glob
import json
import os
import re
import tempfile
from datetime import datetime

import openpyxl

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MANUAL_DIR = os.path.join(REPO_ROOT, "manual")

MONTHS = {"JAN": 1, "FEB": 2, "MAR": 3, "APR": 4, "MAY": 5, "JUN": 6,
          "JUL": 7, "AUG": 8, "SEP": 9, "OCT": 10, "NOV": 11, "DEC": 12}
OWNERS = {"NIC": "Nic", "SHARED": "Shared", "YX": "Yx"}


def find_workbook():
    for candidate in (os.path.join(REPO_ROOT, "Finances.xlsx"),
                      os.path.join(REPO_ROOT, "Archive", "Finances.xlsx")):
        if os.path.exists(candidate):
            return candidate
    found = [p for p in glob.glob(os.path.join(REPO_ROOT, "**", "Finances.xlsx"), recursive=True)
             if os.path.basename(p)[0] != "~"]
    if found:
        return max(found, key=os.path.getmtime)
    raise SystemExit("Could not find Finances.xlsx under " + REPO_ROOT)


def sheet_month(name):
    if "TRIP" in name.upper():
        return None
    tokens = re.split(r"\s+", name.strip())
    if "CC" not in [t.upper() for t in tokens]:
        return None
    month = year = None
    for t in tokens:
        tu = t.upper()
        if tu in MONTHS:
            month = MONTHS[tu]
        elif re.fullmatch(r"\d{2}", t):
            year = 2000 + int(t)
    return ("%04d-%02d" % (year, month)) if (month and year) else None


def clean_description(desc):
    desc = re.sub(r"\s*Ref\s*No\b.*$", "", desc, flags=re.I | re.S)
    return " ".join(desc.split())


def parse_amount(value):
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    if isinstance(value, str):
        s = value.replace(",", "").replace("CR", "").strip()
        try:
            return round(float(s), 2)
        except ValueError:
            return None
    return None


# Same shape build_data.py uses, so a tag written here matches at build time.
# The credit flag keeps a refund from colliding with its original charge. Repeat
# charges still share a key, so each key holds a list of owners in sheet order
# and build_data.py hands them out one per matching row.
def tag_key(month, description, amount, credit):
    return "|".join([month, description.upper()[:60], "%.2f" % amount, "C" if credit else "D"])


def merchant_key(description):
    s = description.upper()
    s = re.sub(r"GPC-[0-9A-Z]+", "", s)
    s = re.sub(r"[0-9]{4,}", "", s)
    s = re.sub(r"[^A-Z ]+", " ", s)
    return " ".join(s.split())[:26]


def load_existing(name):
    path = os.path.join(MANUAL_DIR, name)
    if os.path.exists(path):
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    return {}


def write(name, payload):
    # Write through a temp file: truncating in place meant a crash mid-write
    # destroyed the only copy of a manual file.
    path = os.path.join(MANUAL_DIR, name)
    descriptor, tmp_path = tempfile.mkstemp(
        prefix=name + ".", suffix=".tmp", dir=MANUAL_DIR)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as f:
            json.dump(payload, f, indent=1, ensure_ascii=False)
        os.replace(tmp_path, path)
    finally:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)
    print("wrote", os.path.relpath(path, REPO_ROOT))


def main():
    os.makedirs(MANUAL_DIR, exist_ok=True)
    xlsx = find_workbook()
    print("reading", os.path.relpath(xlsx, REPO_ROOT))
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)

    tags = {}
    merchant_votes = {}
    legacy = []
    for name in wb.sheetnames:
        month = sheet_month(name)
        if not month:
            continue
        for row in wb[name].iter_rows(values_only=True):
            row = list(row) + [None] * 5
            trans, desc, amt_raw, owner_raw = row[1], row[2], row[3], row[4]
            if not isinstance(desc, str):
                continue
            desc = clean_description(desc)
            if desc in ("", "Description of Transaction"):
                continue
            amount = parse_amount(amt_raw)
            if amount is None:
                continue
            credit = isinstance(amt_raw, str) and "CR" in amt_raw.upper()
            legacy.append({
                "date": trans.strftime("%Y-%m-%d") if isinstance(trans, datetime) else None,
                "month": month,
                "card": "UOB ONE CARD",
                "description": desc,
                "amount": amount,
                "credit": credit,
                "foreign": None,
            })
            owner = OWNERS.get(owner_raw.strip().upper()) if isinstance(owner_raw, str) else None
            if not owner:
                continue
            tags.setdefault(tag_key(month, desc, amount, credit), []).append(owner)
            votes = merchant_votes.setdefault(merchant_key(desc), {})
            votes[owner] = votes.get(owner, 0) + 1

    rules = {}
    for merchant, votes in merchant_votes.items():
        total = sum(votes.values())
        top = max(votes, key=votes.get)
        # Only assume a default where history is near-unanimous.
        if total >= 3 and votes[top] / total >= 0.8:
            rules[merchant] = top

    steps, years = [], []
    if "Salary" in wb.sheetnames:
        section = "steps"
        for row in wb["Salary"].iter_rows(values_only=True):
            row = list(row) + [None] * 4
            a, b, c, d = row[0], row[1], row[2], row[3]
            if isinstance(a, str) and a.strip() == "Year":
                section = "years"
                continue
            if isinstance(a, datetime) and isinstance(b, (int, float)):
                steps.append({"from": a.strftime("%Y-%m"), "amount": round(float(b), 2),
                              "note": d.strip() if isinstance(d, str) else None})
            elif section == "years" and isinstance(a, (int, float)) and isinstance(b, (int, float)):
                years.append({"year": int(a), "income": round(float(b), 2),
                              "growth": round(float(c), 4) if isinstance(c, (int, float)) else None,
                              "tax": round(float(d), 2) if isinstance(d, (int, float)) else 0.0})
    steps.sort(key=lambda s: s["from"])
    years.sort(key=lambda y: y["year"])

    # Rewrites merge into the existing file so a re-run cannot destroy keys
    # later migrations or the user added (tagsById, confirmed rules). This
    # script used to emit only _comment + its own key and silently deleted
    # everything else, despite the "safe to re-run" promise.
    tag_file = load_existing("owner_tags.json")
    old_tags = tag_file.get("tags", {})
    merged_tags = dict(tags)
    dropped = 0
    for key, value in old_tags.items():
        # Keys gained a D/C suffix; anything in the old 3-part shape can never
        # match a row again, so drop it rather than carry dead entries forward.
        if len(key.split("|")) != 4:
            dropped += 1
            continue
        merged_tags[key] = value if isinstance(value, list) else [value]
    if dropped:
        print("dropped %d owner tag(s) in the old key format" % dropped)
    tag_file["_comment"] = (
        "Owner tags. Key: month|DESCRIPTION|amount|D or C. Each key holds a "
        "list because a merchant can be charged the same amount twice in a "
        "month; build_data.py assigns them one per matching row. Edit freely.")
    tag_file["tags"] = dict(sorted(merged_tags.items()))
    write("owner_tags.json", tag_file)

    rule_file = load_existing("owner_rules.json")
    old_rules = rule_file.get("rules", {})
    merged_rules = dict(rules)
    merged_rules.update(old_rules)
    rule_file["_comment"] = (
        "Fallback owner per merchant, used when a transaction has no exact tag. "
        "Derived from history where one owner held at least 80% of rows.")
    rule_file["rules"] = dict(sorted(merged_rules.items()))
    write("owner_rules.json", rule_file)

    write("legacy_transactions.json", {
        "_comment": "Card rows straight from the spreadsheet. build_data.py uses these only "
                    "for months with no statement PDF or CSV, so the spreadsheet can be retired.",
        "transactions": legacy
    })

    if not os.path.exists(os.path.join(MANUAL_DIR, "salary.json")):
        write("salary.json", {
            "_comment": "Salary steps and annual income/tax. Gross figures, so they cannot be "
                        "derived from bank deposits - keep this up to date by hand.",
            "steps": steps,
            "years": years
        })
    else:
        print("salary.json already exists, left untouched")

    print()
    print("owner tags: %d" % len(merged_tags))
    print("merchant rules: %d" % len(merged_rules))
    print("salary steps: %d, annual rows: %d" % (len(steps), len(years)))


if __name__ == "__main__":
    main()
